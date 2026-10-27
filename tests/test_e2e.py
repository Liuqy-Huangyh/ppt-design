"""T4 端到端验收：上传案例 → 风格档案 → 生成（pptx + html）→ 反馈 → 档案变化 → 再生成。

只用标准库 unittest（外加工作区自带的 python-pptx 校验产物），在**同进程线程**里
起 `server.app.create_server(host="127.0.0.1", port=0, store_dir=...)`：
不 spawn 子进程、不联网、不需要 LLM key（显式移除 PPT_LLM_* 环境变量）。

运行（仓库根目录）：
    & "$env:USERPROFILE\\.dsh\\dsh-runtimes\\dsh-primary-runtime\\dependencies\\python\\python.exe" tests\\test_e2e.py

覆盖（对应 docs/CONTRACT.md §5/§7 与 T4 验收口径）：
    01  happy path 全链路
        GET  /                      → 200 text/html（浏览器入口页面）
        GET  /api/health            → status=ok、llm=false（离线证明）
        GET  /api/themes            → 至少 3 个主题，含 academic
        POST /api/profiles          → 上传 tests/fixtures/case-academic.pptx，落盘 cases/<id>/source.pptx
        GET  /api/profiles          → 新档案出现在摘要列表
        POST /api/generate          → profile_id + prompt + pages=8，返回 pages/files
        GET  /files/<id>/deck.pptx  → PK 头、Content-Type 正确、python-pptx 能打开、页数与 pages 一致
        GET  /files/<id>/deck.html  → 200、非空、.slide 段落数与 pages 一致
        POST /api/decks/<id>/feedback rating=2 → density 变小、title_scale 变小
        GET  /api/profiles/<id>     → style 已持久化、learn.rounds+1、history 追加
        再次 POST /api/generate     → 用反馈后的档案再生成成功，pptx 仍可打开
    02  缺 outline/prompt → 400 missing_content
    03  上传校验：.txt / 缺 X-Filename / 路径型文件名 → 400 invalid_filename；伪 pptx → 400 invalid_pptx
    04  超过 25MB 的请求 → 413 payload_too_large
    05  路径遍历（/files/../../etc/passwd 等）→ 404 not_found，且不泄漏文件内容
    06  未知 id → 404（profile_not_found / deck_not_found）；非法 rating → 400 invalid_rating
    99  全程 < 60s

临时目录放在工作区 `.tmp/e2e/`（沙箱不允许写系统 temp），测试结束自动清理。
失败信息一律带「方法 + 路径 + 期望值 + 实际响应」，便于定位到具体接口。
"""
from __future__ import annotations

import contextlib
import http.client
import io
import json
import os
import shutil
import sys
import threading
import time
import unittest
import urllib.error
import urllib.parse
import urllib.request
import uuid
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
for extra in (ROOT, ROOT / "kit"):
    if str(extra) not in sys.path:
        sys.path.insert(0, str(extra))

from pptx import Presentation                      # noqa: E402  工作区自带，用于校验产物

from server.app import create_server               # noqa: E402

CASE_PPTX = ROOT / "tests" / "fixtures" / "case-academic.pptx"
WORK_ROOT = ROOT / ".tmp" / "e2e"
LLM_ENV_KEYS = ("PPT_LLM_BASE_URL", "PPT_LLM_API_KEY", "PPT_LLM_MODEL", "PPT_LLM_TIMEOUT")

_MISSING = object()
TIME_BUDGET_SECONDS = 60.0


def _relax_stdio() -> None:
    """Windows 控制台默认 GBK：断言信息里可能出现 ▪ 等字符，避免打印时二次抛异常。"""
    for stream in (sys.stdout, sys.stderr):
        try:
            stream.reconfigure(errors="replace")
        except (AttributeError, ValueError):
            pass


def safe_print(text: str) -> None:
    try:
        print(text, flush=True)
    except UnicodeEncodeError:                      # 极老的 GBK 控制台兜底
        buffer = getattr(sys.stdout, "buffer", None)
        if buffer is not None:
            buffer.write(text.encode("utf-8", "replace") + b"\n")
            buffer.flush()


_relax_stdio()


class E2ETest(unittest.TestCase):
    """案例驱动 + 反馈迭代闭环的端到端验收。"""

    maxDiff = 4000
    started_at = 0.0
    timings: list[tuple[str, float]] = []

    # -- 生命周期 ---------------------------------------------------------- #

    @classmethod
    def setUpClass(cls):
        cls.started_at = time.perf_counter()
        cls.timings = []
        cls.profile_id = None
        cls.deck_ids = []

        # 离线保证：显式移除 LLM 配置；无论 setUpClass 是否成功，结束时都还原。
        cls._saved_env = {key: os.environ.pop(key) for key in LLM_ENV_KEYS
                          if key in os.environ}
        cls.addClassCleanup(cls._restore_env)

        # 临时目录放工作区（系统 temp 在沙箱里不可写）；先注册清理，再起服务。
        cls.work_dir = WORK_ROOT / f"run-{uuid.uuid4().hex[:8]}"
        cls.store_dir = cls.work_dir / "store"
        cls.store_dir.mkdir(parents=True, exist_ok=True)
        cls.addClassCleanup(shutil.rmtree, cls.work_dir, ignore_errors=True)

        cls.httpd, cls.base_url = create_server("127.0.0.1", 0, cls.store_dir)
        cls.port = cls.httpd.server_address[1]
        cls.thread = threading.Thread(target=cls.httpd.serve_forever,
                                      name="ppt-e2e-server", daemon=True)
        cls.thread.start()
        cls.addClassCleanup(cls._stop_server)

    @classmethod
    def _stop_server(cls):
        cls.httpd.shutdown()
        cls.httpd.server_close()
        cls.thread.join(timeout=5)

    @classmethod
    def _restore_env(cls):
        os.environ.update(cls._saved_env)

    @classmethod
    def tearDownClass(cls):
        cls._report_timings()

    @classmethod
    def _report_timings(cls):
        total = time.perf_counter() - cls.started_at
        lines = ["", "E2E 分步耗时（秒）："]
        for label, seconds in cls.timings:
            lines.append(f"  {label:<32} {seconds:6.3f}")
        lines.append(f"  {'TOTAL (incl. setup+teardown)':<30} {total:6.3f}")
        safe_print("\n".join(lines))

    # -- 请求 helpers ------------------------------------------------------ #

    @contextlib.contextmanager
    def timed(self, label: str):
        started = time.perf_counter()
        try:
            yield
        finally:
            type(self).timings.append((label, time.perf_counter() - started))

    def request(self, method: str, path: str, data: bytes | None = None,
                headers: dict | None = None, timeout: float = 120):
        """发请求；HTTP 错误状态也正常返回 (status, body, headers)。"""
        url = self.base_url + path
        req = urllib.request.Request(url, data=data, method=method,
                                     headers=headers or {})
        try:
            with urllib.request.urlopen(req, timeout=timeout) as response:
                return response.status, response.read(), dict(response.headers)
        except urllib.error.HTTPError as exc:
            return exc.code, exc.read(), dict(exc.headers)
        except urllib.error.URLError as exc:
            self.fail(f"{method} {path} 连接失败（服务 {self.base_url} 是否在运行？）：{exc}")

    def raw_request(self, method: str, raw_path: str, body: bytes = b"",
                    headers: dict | None = None, timeout: float = 15):
        """用 http.client 发送**原样**路径（urllib 之外再确认 .. 不被客户端规范化）。"""
        connection = http.client.HTTPConnection("127.0.0.1", self.port, timeout=timeout)
        try:
            connection.putrequest(method, raw_path, skip_accept_encoding=True)
            for key, value in (headers or {}).items():
                connection.putheader(key, value)
            if body:
                connection.putheader("Content-Length", str(len(body)))
            connection.endheaders()
            if body:
                connection.send(body)
            response = connection.getresponse()
            return response.status, response.read()
        finally:
            connection.close()

    def json_call(self, method: str, path: str, payload=_MISSING,
                  headers: dict | None = None, timeout: float = 120):
        head = {"Content-Type": "application/json; charset=utf-8",
                "Accept": "application/json"}
        head.update(headers or {})
        data = None if payload is _MISSING else json.dumps(
            payload, ensure_ascii=False).encode("utf-8")
        status, raw, response_headers = self.request(method, path, data, head,
                                                     timeout=timeout)
        try:
            decoded = json.loads(raw.decode("utf-8"))
        except (UnicodeDecodeError, json.JSONDecodeError):
            self.fail(f"{method} {path} 期望 JSON 响应，实际 HTTP {status}：{raw[:400]!r}")
        return status, decoded, response_headers

    def ok(self, method: str, path: str, payload=_MISSING, headers: dict | None = None):
        status, decoded, _ = self.json_call(method, path, payload, headers)
        self.assertEqual(
            status, 200,
            f"{method} {path} 期望 HTTP 200，实际 {status}；"
            f"响应体：{json.dumps(decoded, ensure_ascii=False)[:800]}")
        return decoded

    def assert_api_error(self, method: str, path: str, status: int, payload,
                         expected_status: int, allowed_codes: set[str]):
        """错误体必须是 {"error":{"code","message"}}，且 code 精确匹配。"""
        actual = json.dumps(payload, ensure_ascii=False)[:800]
        self.assertEqual(status, expected_status,
                         f"{method} {path} 期望 HTTP {expected_status}，实际 {status}；"
                         f"响应体：{actual}")
        self.assertIsInstance(payload, dict,
                              f"{method} {path} 错误响应不是 JSON 对象：{actual}")
        self.assertEqual(set(payload), {"error"},
                         f"{method} {path} 错误体应只有 error 字段，实际：{actual}")
        error = payload["error"]
        self.assertIsInstance(error, dict, f"{method} {path} error 不是对象：{actual}")
        self.assertEqual(set(error), {"code", "message"},
                         f"{method} {path} error 应含 code/message，实际：{actual}")
        self.assertIn(error["code"], allowed_codes,
                      f"{method} {path} 期望 code ∈ {sorted(allowed_codes)}，"
                      f"实际 {error['code']!r}（message={error['message']!r}）")
        self.assertTrue(error["message"],
                        f"{method} {path} 错误 message 不应为空：{actual}")

    def count_profiles(self) -> int:
        return len(self.ok("GET", "/api/profiles"))

    # -- 01：主链路 -------------------------------------------------------- #

    def test_01_full_chain_case_to_feedback_loop(self):
        self.assertTrue(CASE_PPTX.is_file(), f"缺少验收样例：{CASE_PPTX}")

        # 0. 静态首页可访问（浏览器入口；前端细节归 tests/test_web.py）
        with self.timed("GET / (web/index.html)"):
            status, page, page_headers = self.request("GET", "/")
        self.assertEqual(status, 200, f"GET / 期望 200，实际 {status}：{page[:200]!r}")
        self.assertIn("text/html", page_headers.get("Content-Type", ""),
                      f"GET / 的 Content-Type 应为 text/html，实际 "
                      f"{page_headers.get('Content-Type')!r}")
        self.assertIn(b"PPT Studio", page,
                      f"GET / 应返回前端页面（含 PPT Studio 标识），实际 {page[:200]!r}")

        # 1. 健康检查：证明离线（llm=false）且服务就绪
        with self.timed("GET /api/health"):
            health = self.ok("GET", "/api/health")
        self.assertEqual(health.get("status"), "ok",
                         f"GET /api/health 期望 status=ok，实际 {health}")
        self.assertIs(health.get("llm"), False,
                      "GET /api/health 期望 llm=false（离线验收），实际 "
                      f"{health.get('llm')!r}；测试环境不应配置 PPT_LLM_*")
        self.assertEqual(health.get("version"), "1.0.0",
                         f"GET /api/health 版本应为 1.0.0，实际 {health.get('version')!r}")

        # 2. 主题列表：档案套用的基底主题必须存在
        with self.timed("GET /api/themes"):
            themes = self.ok("GET", "/api/themes")
        self.assertIsInstance(themes, list, f"GET /api/themes 应返回数组，实际 {themes!r}")
        keys = {item.get("key") for item in themes}
        self.assertIn("academic", keys,
                      f"GET /api/themes 应包含 academic，实际 {sorted(keys)}")
        self.assertGreaterEqual(len(themes), 3,
                                f"GET /api/themes 至少 3 个主题，实际 {keys}")
        for item in themes:
            self.assertEqual(set(item), {"key", "label"},
                             f"GET /api/themes 每项应为 {{key,label}}，实际 {item}")

        # 3. 上传案例 → 风格档案
        source_prs = Presentation(str(CASE_PPTX))
        source_slides = len(source_prs.slides)
        source_size = [round(float(source_prs.slide_width) / 914400, 2),
                       round(float(source_prs.slide_height) / 914400, 2)]
        raw = CASE_PPTX.read_bytes()
        with self.timed("POST /api/profiles"):
            status, body, headers = self.request(
                "POST", "/api/profiles", data=raw,
                headers={"Content-Type": "application/octet-stream",
                         "X-Filename": urllib.parse.quote(CASE_PPTX.name)})
        self.assertEqual(status, 200,
                         f"POST /api/profiles 期望 200，实际 {status}：{body[:600]!r}")
        self.assertIn("application/json", headers.get("Content-Type", ""),
                      "POST /api/profiles 应返回 JSON")
        uploaded = json.loads(body.decode("utf-8"))
        self.assertEqual(set(uploaded), {"id", "profile"},
                         f"POST /api/profiles 应返回 {{id,profile}}，实际 {list(uploaded)}")

        profile = uploaded["profile"]
        profile_id = uploaded["id"]
        self.assertTrue(profile_id.startswith("pf-"),
                        f"POST /api/profiles 返回的档案 id 应以 pf- 开头，实际 {profile_id!r}")
        self.assertEqual(profile.get("id"), profile_id,
                         "POST /api/profiles 的 id 与 profile.id 应一致")
        for key in ("id", "name", "source", "created_at", "slide_size",
                    "palette", "fonts", "metrics", "style", "learn"):
            self.assertIn(key, profile,
                          f"POST /api/profiles 返回的档案缺少字段 {key}："
                          f"{json.dumps(profile, ensure_ascii=False)[:600]}")
        self.assertEqual(profile["name"], CASE_PPTX.stem,
                         f"档案名应取上传文件名（{CASE_PPTX.stem}），实际 {profile['name']!r}")
        self.assertEqual(profile["source"], "source.pptx",
                         "档案 source 应为 source.pptx（不暴露原始路径）")
        self.assertEqual(profile["metrics"].get("slides"), source_slides,
                         "档案 metrics.slides 应等于案例实际页数 "
                         f"({source_slides})，实际 {profile['metrics'].get('slides')!r}")
        self.assertEqual([round(float(v), 2) for v in profile["slide_size"]],
                         source_size,
                         "档案 slide_size 应等于案例幻灯片尺寸（英寸） "
                         f"{source_size}，实际 {profile['slide_size']!r}")
        self.assertEqual(profile["style"].get("density"), 1.0,
                         "新档案的 style.density 基准应为 1.0，实际 "
                         f"{profile['style'].get('density')!r}")
        self.assertEqual(profile["learn"].get("rounds"), 0,
                         f"新档案 learn.rounds 应为 0，实际 {profile['learn'].get('rounds')!r}")
        # 案例分析真的读出了案例配色（只读 fixture case-academic 的实测值）
        self.assertEqual(profile["palette"].get("accent"), "#1B3A6B",
                         "档案 palette.accent 应来自案例主题色 #1B3A6B，实际 "
                         f"{profile['palette'].get('accent')!r}")
        for key, value in (profile["palette"] or {}).items():
            self.assertRegex(str(value), r"^#[0-9A-F]{6}$",
                             f"档案 palette.{key} 应为大写 #RRGGBB，实际 {value!r}")

        case_copy = self.store_dir / "cases" / profile_id / "source.pptx"
        self.assertTrue(case_copy.is_file(),
                        f"上传的案例应落盘到 cases/{profile_id}/source.pptx")
        self.assertEqual(case_copy.read_bytes(), raw,
                         "cases/<id>/source.pptx 应与上传字节完全一致")

        type(self).profile_id = profile_id

        # 4. 档案出现在列表里
        with self.timed("GET /api/profiles"):
            summaries = self.ok("GET", "/api/profiles")
        ids = [item.get("id") for item in summaries]
        self.assertIn(profile_id, ids,
                      f"GET /api/profiles 摘要列表应包含 {profile_id}，实际 {ids}")
        summary = next(item for item in summaries if item.get("id") == profile_id)
        self.assertIn("palette", summary, f"档案摘要应含 palette，实际 {summary}")
        self.assertEqual(summary.get("learn"), {"rounds": 0},
                         f"档案摘要 learn.rounds 应为 0，实际 {summary.get('learn')!r}")

        # 5. 用该档案 + 提示词生成 deck（离线大纲）
        pages_requested = 8
        with self.timed("POST /api/generate #1"):
            deck = self.ok("POST", "/api/generate", {
                "title": "案例驱动验收",
                "prompt": "按照上传的学术答辩案例风格，讲清楚案例驱动 + 反馈迭代的闭环",
                "profile_id": profile_id,
                "pages": pages_requested,
                "theme": "academic",
            })
        for key in ("id", "pages", "files", "notes", "llm_used", "profile_id", "theme"):
            self.assertIn(key, deck,
                          f"POST /api/generate 响应缺少字段 {key}："
                          f"{json.dumps(deck, ensure_ascii=False)[:600]}")
        self.assertTrue(str(deck.get("id", "")).startswith("dk-"),
                        f"POST /api/generate 返回 id 应以 dk- 开头，实际 {deck.get('id')!r}")
        self.assertEqual(deck.get("profile_id"), profile_id,
                         "POST /api/generate 应回显 profile_id")
        self.assertEqual(deck.get("theme"), "academic",
                         f"theme 应为 academic，实际 {deck.get('theme')!r}")
        self.assertIs(deck.get("llm_used"), False,
                      "离线生成时 llm_used 应为 false，实际 "
                      f"{deck.get('llm_used')!r}（测试环境不应配置 LLM key）")
        pages = deck.get("pages")
        self.assertIsInstance(pages, int,
                              f"POST /api/generate 的 pages 应为整数，实际 {pages!r}")
        self.assertGreater(pages, 0,
                           f"POST /api/generate 的 pages 应 > 0，实际 {pages}")
        self.assertGreaterEqual(pages, 3,
                                f"请求 {pages_requested} 页，实际只有 {pages} 页（离线大纲过短）")
        expected_files = {"pptx": f"/files/{deck['id']}/deck.pptx",
                          "html": f"/files/{deck['id']}/deck.html"}
        self.assertEqual(deck.get("files"), expected_files,
                         f"files 下载地址不对，期望 {expected_files}，实际 {deck.get('files')!r}")
        notes = deck.get("notes") or []
        self.assertTrue(any("档案" in str(note) for note in notes),
                        f"notes 应说明用了哪个档案，实际 {notes}")
        self.assertTrue(any(("离线" in str(note)) or ("LLM" in str(note)) for note in notes),
                        f"notes 应说明离线大纲/LLM 状态，实际 {notes}")

        deck_id = deck["id"]
        type(self).deck_ids.append(deck_id)

        # 6. 下载 pptx：可被 python-pptx 打开，页数与接口返回一致
        with self.timed("GET deck.pptx #1"):
            status, pptx_bytes, pptx_headers = self.request(
                "GET", deck["files"]["pptx"])
        self.assertEqual(status, 200,
                         f"GET {deck['files']['pptx']} 期望 200，实际 {status}："
                         f"{pptx_bytes[:300]!r}")
        self.assertTrue(pptx_bytes.startswith(b"PK"),
                        "下载的 deck.pptx 应以 ZIP 魔数 PK 开头，实际 "
                        f"{pptx_bytes[:8]!r}")
        self.assertIn("presentationml",
                      pptx_headers.get("Content-Type", ""),
                      "deck.pptx 的 Content-Type 应为 presentationml，实际 "
                      f"{pptx_headers.get('Content-Type')!r}")
        try:
            prs = Presentation(io.BytesIO(pptx_bytes))
        except Exception as exc:                     # noqa: BLE001 - 失败即验收失败
            self.fail(f"GET {deck['files']['pptx']} 下载的 pptx 无法用 python-pptx 打开："
                      f"{type(exc).__name__}: {exc}")
        self.assertEqual(len(prs.slides), pages,
                         f"下载的 deck.pptx 实际 {len(prs.slides)} 页，"
                         f"与 /api/generate 返回的 pages={pages} 不一致")
        self.assertGreater(int(prs.slide_width or 0), 0,
                           "生成的 deck.pptx 幻灯片宽度应大于 0（否则是坏文件）")
        disk_pptx = self.store_dir / "decks" / deck_id / "deck.pptx"
        self.assertTrue(disk_pptx.is_file(),
                        f"deck 落盘文件缺失：decks/{deck_id}/deck.pptx")
        self.assertEqual(disk_pptx.read_bytes(), pptx_bytes,
                         "GET /files/.../deck.pptx 的字节应与落盘文件一致")

        # 7. 下载网页版：非空、页数一致（1280×720 舞台与 iframe 缩放是前端已知行为，不做视觉断言）
        with self.timed("GET deck.html #1"):
            status, html_bytes, html_headers = self.request(
                "GET", deck["files"]["html"])
        self.assertEqual(status, 200,
                         f"GET {deck['files']['html']} 期望 200，实际 {status}")
        self.assertIn("text/html", html_headers.get("Content-Type", ""),
                      "deck.html 的 Content-Type 应为 text/html，实际 "
                      f"{html_headers.get('Content-Type')!r}")
        self.assertGreater(len(html_bytes), 2000,
                           f"deck.html 过短（{len(html_bytes)} 字节），疑似空页")
        html_text = html_bytes.decode("utf-8", "replace")
        self.assertIn("<!doctype html", html_text.lower(),
                      "deck.html 应是完整 HTML 文档（缺少 <!DOCTYPE html>）")
        slide_sections = html_text.count('class="slide')
        self.assertEqual(slide_sections, pages,
                         f"deck.html 的 .slide 段落数（{slide_sections}）"
                         f"应与 pages（{pages}）一致")

        # 8. 评分反馈 rating=2 → 档案 density 下降、rounds+1
        before = self.ok("GET", f"/api/profiles/{profile_id}")
        density_before = float(before["style"]["density"])
        rounds_before = int(before["learn"]["rounds"])
        with self.timed("POST feedback rating=2"):
            feedback = self.ok("POST", f"/api/decks/{deck_id}/feedback",
                               {"rating": 2, "notes": "要点太密，标题也偏大"})
        self.assertIs(feedback.get("ok"), True,
                      f"POST /api/decks/{deck_id}/feedback 应返回 ok=true，实际 {feedback}")
        self.assertEqual(feedback.get("profile_id"), profile_id,
                         "反馈响应应指向被调整的档案")
        style = feedback.get("style") or {}
        self.assertLess(float(style.get("density", 9)),
                        density_before,
                        "rating=2 后 density 应下降：反馈前 "
                        f"{density_before}，反馈后 {style.get('density')!r}")
        self.assertLess(float(style.get("title_scale", 9)), 1.0,
                        "rating=2 后 title_scale 应小于 1.0，实际 "
                        f"{style.get('title_scale')!r}")
        self.assertTrue(str(feedback.get("message", "")).startswith("已调整"),
                        f"反馈 message 应以「已调整」开头，实际 {feedback.get('message')!r}")

        # 9. 档案已持久化：GET 回来的参数与反馈一致，history 追加
        with self.timed("GET /api/profiles/<id>"):
            after = self.ok("GET", f"/api/profiles/{profile_id}")
        self.assertAlmostEqual(float(after["style"]["density"]),
                               float(style["density"]), places=6,
                               msg="反馈后的 density 应立即持久化到档案文件")
        self.assertEqual(int(after["learn"]["rounds"]), rounds_before + 1,
                         f"learn.rounds 应 +1：反馈前 {rounds_before}，"
                         f"反馈后 {after['learn'].get('rounds')!r}")
        history = after["learn"].get("history") or []
        self.assertTrue(history, "learn.history 应追加一条记录")
        self.assertEqual(history[-1].get("rating"), 2,
                         f"history 最后一条 rating 应为 2，实际 {history[-1]!r}")

        feedback_log = self.store_dir / "feedback.jsonl"
        self.assertTrue(feedback_log.is_file(), "反馈应追加落盘到 feedback.jsonl")
        lines = [line for line in feedback_log.read_text(encoding="utf-8").splitlines()
                 if line.strip()]
        self.assertTrue(lines, "feedback.jsonl 不应为空")
        record = json.loads(lines[-1])
        self.assertEqual(record.get("deck_id"), deck_id,
                         f"feedback.jsonl 最后一条 deck_id 应为 {deck_id}，实际 {record}")
        self.assertEqual(record.get("profile_id"), profile_id,
                         f"feedback.jsonl 最后一条 profile_id 应为 {profile_id}")
        self.assertEqual(record.get("rating"), 2,
                         f"feedback.jsonl 最后一条 rating 应为 2，实际 {record.get('rating')!r}")

        # 10. 用反馈后的档案再生成一次（不要求像素级差异，只要求链路成立）
        with self.timed("POST /api/generate #2"):
            deck2 = self.ok("POST", "/api/generate", {
                "title": "反馈后的再生成",
                "prompt": "用调整后的档案再生成一版，验证反馈闭环可用",
                "profile_id": profile_id,
                "pages": 8,
                "theme": "academic",
            })
        self.assertEqual(deck2.get("profile_id"), profile_id,
                         "第二次生成仍应使用同一档案")
        self.assertGreater(deck2.get("pages", 0), 0,
                           f"第二次生成的 pages 应 > 0，实际 {deck2.get('pages')!r}")
        self.assertNotEqual(deck2.get("id"), deck_id,
                            "第二次生成应产出新的 deck id")
        type(self).deck_ids.append(deck2["id"])
        with self.timed("GET deck.pptx #2"):
            status, pptx2, _ = self.request("GET", deck2["files"]["pptx"])
        self.assertEqual(status, 200,
                         f"GET {deck2['files']['pptx']} 期望 200，实际 {status}")
        try:
            prs2 = Presentation(io.BytesIO(pptx2))
        except Exception as exc:                     # noqa: BLE001
            self.fail(f"GET {deck2['files']['pptx']} 的 pptx 无法打开："
                      f"{type(exc).__name__}: {exc}")
        self.assertEqual(len(prs2.slides), deck2["pages"],
                         "第二次生成的 deck.pptx 页数与 pages 不一致")

        # 11. deck 列表 / 元数据能对上（前端「最近生成」依赖这两个接口）
        with self.timed("GET /api/decks"):
            decks = self.ok("GET", "/api/decks")
        listed = {item.get("id"): item for item in decks}
        for did in (deck_id, deck2["id"]):
            self.assertIn(did, listed,
                          f"GET /api/decks 应包含 {did}，实际 {list(listed)}")
            self.assertEqual(listed[did].get("profile_id"), profile_id,
                             f"deck {did} 的摘要 profile_id 应为 {profile_id}")
        meta = self.ok("GET", f"/api/decks/{deck_id}")
        self.assertEqual(meta.get("pages"), pages,
                         f"GET /api/decks/{deck_id} 的 pages 应为 {pages}，"
                         f"实际 {meta.get('pages')!r}")
        self.assertEqual(meta.get("profile_id"), profile_id,
                         f"meta.json 的 profile_id 应为 {profile_id}")
        self.assertIs(meta.get("llm_used"), False,
                      f"meta.json 的 llm_used 应为 false，实际 {meta.get('llm_used')!r}")

    # -- 02：缺内容 -------------------------------------------------------- #

    def test_02_generate_without_content_is_400(self):
        for payload in ({}, {"title": "只有标题"}, {"prompt": "   "},
                        {"outline": "", "prompt": ""}):
            with self.subTest(payload=payload):
                status, decoded, _ = self.json_call("POST", "/api/generate", payload)
                self.assert_api_error("POST", "/api/generate", status, decoded,
                                      400, {"missing_content"})

    # -- 03：上传校验 ------------------------------------------------------ #

    def test_03_upload_rejects_bad_files(self):
        before = self.count_profiles()
        raw = CASE_PPTX.read_bytes()
        cases = [
            ("扩展名 .txt",
             b"not a pptx at all",
             {"X-Filename": urllib.parse.quote("案例.txt")},
             400, {"invalid_filename"}),
            ("缺少 X-Filename 头",
             b"not a pptx at all",
             {},
             400, {"invalid_filename"}),
            ("文件名含路径 ../",
             b"not a pptx at all",
             {"X-Filename": "../../evil.pptx"},
             400, {"invalid_filename"}),
            ("内容是垃圾字节",
             b"this is definitely not a pptx package",
             {"X-Filename": "fake.pptx"},
             400, {"invalid_pptx"}),
            ("旧版 .ppt（OLE 头）",
             b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 512,
             {"X-Filename": "legacy.ppt"},
             400, {"invalid_pptx"}),
        ]
        for label, body, headers, expected_status, allowed_codes in cases:
            with self.subTest(case=label):
                status, raw_body, _ = self.request("POST", "/api/profiles",
                                                   data=body, headers=headers)
                try:
                    decoded = json.loads(raw_body.decode("utf-8"))
                except (UnicodeDecodeError, json.JSONDecodeError):
                    self.fail(f"POST /api/profiles（{label}）期望 JSON 错误体，"
                              f"实际 HTTP {status}：{raw_body[:300]!r}")
                self.assert_api_error("POST", "/api/profiles", status, decoded,
                                      expected_status, allowed_codes)

        # 合法 pptx 但缺 X-Filename 也必须被拒（扩展名来自头部，不猜）
        status, raw_body, _ = self.request("POST", "/api/profiles", data=raw)
        self.assert_api_error("POST", "/api/profiles", status,
                              json.loads(raw_body.decode("utf-8")),
                              400, {"invalid_filename"})

        self.assertEqual(self.count_profiles(), before,
                         "被拒绝的上传不应产生档案（GET /api/profiles 数量变多了）")

    # -- 04：超大请求 ------------------------------------------------------ #

    def test_04_oversize_upload_is_413(self):
        """25 MB 上限由 Content-Length 先拦下（服务端不读取超限请求体）。

        真发 26 MB 时 Windows 会在服务端提前关连接时复位套接字，客户端拿不到响应；
        这里按契约构造「声明 26 MB 的请求」，验证拦截发生在读取请求体之前。
        """
        before = self.count_profiles()
        declared = 26 * 1024 * 1024
        connection = http.client.HTTPConnection("127.0.0.1", self.port, timeout=15)
        try:
            connection.putrequest("POST", "/api/profiles")
            connection.putheader("X-Filename", "big.pptx")
            connection.putheader("Content-Type", "application/octet-stream")
            connection.putheader("Content-Length", str(declared))
            connection.endheaders()
            response = connection.getresponse()
            status = response.status
            body = response.read()
        finally:
            connection.close()
        try:
            decoded = json.loads(body.decode("utf-8"))
        except (UnicodeDecodeError, json.JSONDecodeError):
            self.fail(f"POST /api/profiles（声明 {declared} 字节）期望 JSON 错误体，"
                      f"实际 HTTP {status}：{body[:300]!r}")
        self.assert_api_error("POST", "/api/profiles", status, decoded,
                              413, {"payload_too_large"})
        self.assertEqual(self.count_profiles(), before,
                         "413 的超大上传不应产生档案")

        # JSON 接口同样有 1 MB 上限：声明 2 MB 也应 413
        connection = http.client.HTTPConnection("127.0.0.1", self.port, timeout=15)
        try:
            connection.putrequest("POST", "/api/generate")
            connection.putheader("Content-Type", "application/json")
            connection.putheader("Content-Length", str(2 * 1024 * 1024))
            connection.endheaders()
            response = connection.getresponse()
            status = response.status
            body = response.read()
        finally:
            connection.close()
        decoded = json.loads(body.decode("utf-8"))
        self.assert_api_error("POST", "/api/generate", status, decoded,
                              413, {"payload_too_large"})

    # -- 05：路径遍历 ------------------------------------------------------ #

    def test_05_path_traversal_is_404(self):
        deck_id = self.deck_ids[0] if self.deck_ids else "dk-20000101-dead"
        paths = [
            "/files/" + deck_id + "/../../etc/passwd",
            "/files/" + deck_id + "/..%2F..%2Fmeta.json",
            "/files/" + deck_id + "/%2e%2e%2fdeck.html",
            "/api/../health",
            "/../server/app.py",
        ]
        for path in paths:
            with self.subTest(path=path):
                status, body = self.raw_request("GET", path)
                text = body.decode("utf-8", "replace")
                try:
                    decoded = json.loads(text)
                except json.JSONDecodeError:
                    self.fail(f"GET {path} 期望 JSON 错误体，实际 HTTP {status}："
                              f"{text[:300]!r}")
                self.assert_api_error("GET", path, status, decoded,
                                      404, {"not_found"})
                lowered = text.lower()
                for marker in ("root:x:", "import argparse", "def main("):
                    self.assertNotIn(marker, lowered,
                                     f"GET {path} 泄漏了文件内容（包含 {marker!r}）："
                                     f"{text[:300]!r}")

    # -- 06：未知 id / 非法参数 -------------------------------------------- #

    def test_06_unknown_ids_and_invalid_rating(self):
        status, decoded, _ = self.json_call("GET", "/api/profiles/pf-20000101-dead")
        self.assert_api_error("GET", "/api/profiles/pf-20000101-dead", status, decoded,
                              404, {"profile_not_found"})

        status, decoded, _ = self.json_call("GET", "/api/decks/dk-20000101-dead")
        self.assert_api_error("GET", "/api/decks/dk-20000101-dead", status, decoded,
                              404, {"deck_not_found"})

        status, decoded, _ = self.json_call("POST", "/api/generate",
                                            {"prompt": "任何内容",
                                             "profile_id": "pf-20000101-dead"})
        self.assert_api_error("POST", "/api/generate", status, decoded,
                              404, {"profile_not_found"})

        status, decoded, _ = self.json_call("POST", "/api/generate",
                                            {"prompt": "任何内容",
                                             "theme": "no-such-theme"})
        self.assert_api_error("POST", "/api/generate", status, decoded,
                              400, {"invalid_theme"})

        status, decoded, _ = self.json_call(
            "POST", "/api/decks/dk-20000101-dead/feedback", {"rating": 3})
        self.assert_api_error("POST", "/api/decks/dk-20000101-dead/feedback",
                              status, decoded, 404, {"deck_not_found"})

        if self.deck_ids:
            deck_id = self.deck_ids[0]
            for rating in (0, 6, "5", True, None):
                with self.subTest(rating=rating):
                    status, decoded, _ = self.json_call(
                        "POST", f"/api/decks/{deck_id}/feedback", {"rating": rating})
                    self.assert_api_error("POST", f"/api/decks/{deck_id}/feedback",
                                          status, decoded, 400, {"invalid_rating"})

    # -- 99：时间预算 ------------------------------------------------------ #

    def test_99_within_time_budget(self):
        elapsed = time.perf_counter() - type(self).started_at
        self.assertLess(elapsed, TIME_BUDGET_SECONDS,
                        f"端到端测试耗时 {elapsed:.1f}s，超过 {TIME_BUDGET_SECONDS:.0f}s 预算")


if __name__ == "__main__":
    unittest.main(verbosity=2)
