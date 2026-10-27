"""T2 内核自测：案例分析 / 风格档案 / 反馈调参 / 提示词大纲 / 构建不回归。

跑法（工作区根目录）：
    & "$env:USERPROFILE\\.dsh\\dsh-runtimes\\dsh-primary-runtime\\dependencies\\python\\python.exe" tests\\test_profile.py

注：沙箱不允许写系统临时目录，测试用工作区内的 .tmp/t2-tests/ 当临时区。
"""
from __future__ import annotations

import contextlib
import json
import os
import re
import shutil
import sys
import unittest
import uuid
from pathlib import Path
from unittest import mock

ROOT = Path(__file__).resolve().parents[1]
KIT = ROOT / "kit"
if str(KIT) not in sys.path:
    sys.path.insert(0, str(KIT))

from pptx import Presentation                       # noqa: E402
from pptx.util import Inches                         # noqa: E402

from pptxkit import builder, markdown as md, outline  # noqa: E402
from pptxkit import profile as profiler              # noqa: E402
from pptxkit import themes                           # noqa: E402

TMP_ROOT = ROOT / ".tmp" / "t2-tests"


@contextlib.contextmanager
def temp_dir(name: str = "case"):
    """工作区内的临时目录（系统 temp 在沙箱里不可写）。"""
    TMP_ROOT.mkdir(parents=True, exist_ok=True)
    path = TMP_ROOT / f"{name}-{os.getpid()}-{uuid.uuid4().hex[:8]}"
    path.mkdir(parents=True, exist_ok=True)
    try:
        yield path
    finally:
        shutil.rmtree(path, ignore_errors=True)

SAMPLE = ROOT / "tests" / "fixtures" / "case-academic.pptx"
EDITORIAL = ROOT / "tests" / "fixtures" / "case-editorial.pptx"
BUSINESS_MD = ROOT / "templates" / "business-minimal.md"

HEX_RE = re.compile(r"^#[0-9A-F]{6}$")
ID_RE = re.compile(r"^pf-\d{8}-[0-9a-f]{8,}$")


class AnalyzeTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.profile = profiler.analyze(SAMPLE)

    def test_has_contract_sections(self):
        for key in ("id", "name", "source", "created_at", "slide_size",
                    "palette", "fonts", "metrics", "style", "learn"):
            self.assertIn(key, self.profile)
        self.assertEqual(self.profile["slide_size"], [13.33, 7.5])

    def test_palette_is_rrggbb(self):
        palette = self.profile["palette"]
        self.assertEqual(set(palette), set(profiler.PALETTE_KEYS))
        for key, value in palette.items():
            self.assertTrue(HEX_RE.match(value), f"{key}={value!r} 不是 #RRGGBB")
        self.assertNotEqual(palette["bg"], palette["ink"])
        self.assertNotEqual(palette["accent"], palette["bg"])

    def test_palette_reflects_sample(self):
        """样例是 academic 主题生成的：重点色应被真实读出来，而不是默认值。"""
        palette = self.profile["palette"]
        self.assertEqual(palette["bg"], "#FFFFFF")
        self.assertEqual(palette["ink"], "#14213D")
        self.assertEqual(palette["accent"], "#1B3A6B")
        self.assertEqual(palette["accent_soft"], "#E9EEF7")
        self.assertEqual(palette["rule"], "#D5DCE8")

    def test_fonts_and_metrics(self):
        fonts = self.profile["fonts"]
        self.assertEqual(fonts["title"], "微软雅黑")
        self.assertEqual(fonts["body"], "微软雅黑")

        metrics = self.profile["metrics"]
        self.assertEqual(metrics["slides"], 17)
        self.assertGreater(metrics["chars_per_slide"], 10)
        self.assertGreater(metrics["bullets_per_slide"], 1)
        self.assertTrue(10 <= metrics["title_pt"] <= 80)
        self.assertTrue(8 <= metrics["body_pt"] <= 48)
        self.assertTrue(0 <= metrics["image_ratio"] <= 1)

    def test_editorial_fonts_come_from_runs(self):
        editorial = profiler.analyze(EDITORIAL)
        self.assertEqual(editorial["fonts"]["title"], "宋体")
        self.assertEqual(editorial["fonts"]["body"], "微软雅黑")
        self.assertEqual(editorial["palette"]["accent"], "#1F3BFF")

    def test_corrupt_file_returns_default(self):
        with temp_dir() as tmp:
            bad = Path(tmp) / "broken.pptx"
            bad.write_bytes(b"this is definitely not a pptx")
            result = profiler.analyze(bad)
        self.assertEqual(result["palette"], profiler.DEFAULT_PALETTE)
        self.assertEqual(result["fonts"], profiler.DEFAULT_FONTS)
        self.assertEqual(result["metrics"]["slides"], 0)
        self.assertEqual(result["name"], "broken")

    def test_missing_file_returns_default(self):
        with temp_dir() as tmp:
            result = profiler.analyze(Path(tmp) / "nope.pptx")
        self.assertEqual(result["metrics"]["slides"], 0)
        self.assertTrue(HEX_RE.match(result["palette"]["bg"]))

    def test_empty_pptx_returns_default(self):
        with temp_dir() as tmp:
            empty = Path(tmp) / "empty.pptx"
            prs = Presentation()
            prs.slide_width = Inches(13.333)
            prs.slide_height = Inches(7.5)
            prs.save(str(empty))
            result = profiler.analyze(empty)
        self.assertEqual(result["metrics"]["slides"], 0)
        self.assertEqual(result["palette"], profiler.DEFAULT_PALETTE)
        for key, value in result["palette"].items():
            self.assertTrue(HEX_RE.match(value), f"{key}={value!r}")

    def test_analyze_does_not_mutate_caller_path(self):
        result = profiler.analyze(str(SAMPLE))
        self.assertEqual(result["source"], "case-academic.pptx")
        self.assertIn("_source_path", result)


class StoreTests(unittest.TestCase):
    def setUp(self):
        self._tmp = temp_dir("store")
        self.workdir = self._tmp.__enter__()
        self.store = self.workdir / "store"
        self.profile = profiler.analyze(SAMPLE, name="我的答辩模板")

    def tearDown(self):
        self._tmp.__exit__(None, None, None)

    def test_save_layout_and_roundtrip(self):
        saved = profiler.save(self.profile, self.store)
        pid = saved["id"]
        self.assertTrue(ID_RE.match(pid), pid)

        profile_json = self.store / "profiles" / f"{pid}.json"
        source_pptx = self.store / "cases" / pid / "source.pptx"
        self.assertTrue(profile_json.is_file())
        self.assertTrue(source_pptx.is_file())
        self.assertEqual(saved["source"], "source.pptx")

        on_disk = json.loads(profile_json.read_text(encoding="utf-8"))
        self.assertNotIn("_source_path", on_disk)

        loaded = profiler.load(self.store, pid)
        self.assertEqual(loaded, saved)
        self.assertEqual(loaded["name"], "我的答辩模板")
        self.assertEqual(loaded["palette"]["accent"], "#1B3A6B")

    def test_list_all_sorted_and_complete(self):
        first = profiler.save(self.profile, self.store)
        second_profile = profiler.analyze(EDITORIAL, name="杂志分享")
        second = profiler.save(second_profile, self.store)
        listed = profiler.list_all(self.store)
        self.assertEqual({item["id"] for item in listed}, {first["id"], second["id"]})
        self.assertEqual(len(listed), 2)
        self.assertGreaterEqual(listed[0]["created_at"], listed[1]["created_at"])
        for item in listed:
            self.assertIn("palette", item)
            self.assertTrue(HEX_RE.match(item["palette"]["bg"]))

    def test_list_all_empty_dir(self):
        self.assertEqual(profiler.list_all(self.store / "missing"), [])

    def test_load_rejects_bad_id(self):
        for bad in ("../secret", "..", "a/b", "", "x" * 80):
            with self.assertRaises(FileNotFoundError):
                profiler.load(self.store, bad)

    def test_load_missing_id(self):
        with self.assertRaises(FileNotFoundError):
            profiler.load(self.store, "pf-20260101-abcd")

    def test_save_sanitizes_hostile_id(self):
        profile = dict(self.profile)
        profile["id"] = "../../evil"
        saved = profiler.save(profile, self.store)
        self.assertTrue(ID_RE.match(saved["id"]), saved["id"])
        self.assertTrue((self.store / "profiles" / f"{saved['id']}.json").is_file())
        self.assertFalse((self.store.parent / "evil.json").exists())

    def test_save_is_stable_for_same_id(self):
        saved = profiler.save(self.profile, self.store)
        again = profiler.save(saved, self.store)
        self.assertEqual(again["id"], saved["id"])
        self.assertEqual(len(profiler.list_all(self.store)), 1)
        self.assertTrue((self.store / "cases" / saved["id"] / "source.pptx").is_file())


class ApplyTests(unittest.TestCase):
    def test_apply_returns_theme_with_overrides(self):
        profile = profiler.analyze(SAMPLE)
        theme = profiler.apply(profile)
        self.assertIsInstance(theme, themes.Theme)
        self.assertEqual(theme.bg, themes.rgb("FFFFFF"))
        self.assertEqual(theme.ink, themes.rgb("14213D"))
        self.assertEqual(theme.accent, themes.rgb("1B3A6B"))
        self.assertEqual(theme.font_title, "微软雅黑")
        self.assertEqual(theme.font_body, "微软雅黑")
        base = themes.get_theme(None)
        self.assertEqual(theme.title_size, round(base.title_size * 1.0))

    def test_apply_scales_title_and_keeps_base_layout(self):
        profile = profiler.analyze(SAMPLE)
        profile["style"] = {"density": 1.0, "title_scale": 1.2, "bullet_char": "▪"}
        theme = profiler.apply(profile, "minimal")
        self.assertEqual(theme.key, "minimal")
        self.assertEqual(theme.title_size, round(themes.get_theme("minimal").title_size * 1.2))
        self.assertEqual(theme.body_size, themes.get_theme("minimal").body_size)

    def test_apply_clamps_and_survives_garbage(self):
        profile = profiler.analyze(SAMPLE)
        profile["palette"] = {key: "not-a-color" for key in profiler.PALETTE_KEYS}
        profile["style"] = {"title_scale": 99, "density": 99}
        theme = profiler.apply(profile, "editorial")
        self.assertEqual(theme.title_size,
                         round(themes.get_theme("editorial").title_size * 1.25))
        for color in (theme.bg, theme.ink, theme.muted, theme.accent,
                      theme.accent_soft, theme.rule):
            self.assertEqual(len(str(color)), 6)

    def test_clone_with_keeps_existing_themes(self):
        base = themes.get_theme("academic")
        clone = themes.clone_with(base, accent=themes.rgb("FF0000"))
        self.assertEqual(clone.accent, themes.rgb("FF0000"))
        self.assertEqual(base.accent, themes.rgb("1B3A6B"))     # 原主题不被改动
        for key in ("academic", "editorial", "minimal"):
            self.assertIn(key, themes.THEMES)


class LearnTests(unittest.TestCase):
    def _fresh(self):
        return profiler.analyze(SAMPLE)

    def test_high_rating_raises_density(self):
        profile = self._fresh()
        result = profiler.learn(profile, 5, "很好")
        self.assertIs(result, profile)
        self.assertAlmostEqual(profile["style"]["density"], 1.03, places=4)
        self.assertAlmostEqual(profile["style"]["title_scale"], 1.0, places=4)
        self.assertEqual(profile["learn"]["rounds"], 1)
        entry = profile["learn"]["history"][-1]
        self.assertEqual(set(entry), {"at", "rating", "density", "title_scale", "notes"})
        self.assertEqual(entry["rating"], 5)
        self.assertEqual(entry["notes"], "很好")

    def test_middle_rating_changes_nothing(self):
        profile = self._fresh()
        profiler.learn(profile, 3)
        self.assertAlmostEqual(profile["style"]["density"], 1.0, places=4)
        self.assertAlmostEqual(profile["style"]["title_scale"], 1.0, places=4)
        self.assertEqual(profile["learn"]["rounds"], 1)
        self.assertEqual(len(profile["learn"]["history"]), 1)

    def test_low_rating_reduces_density_and_title(self):
        profile = self._fresh()
        profiler.learn(profile, 1, "太挤了")
        self.assertAlmostEqual(profile["style"]["density"], 0.9, places=4)
        self.assertAlmostEqual(profile["style"]["title_scale"], 0.97, places=4)

    def test_clamped_to_contract_ranges(self):
        profile = self._fresh()
        for _ in range(30):
            profiler.learn(profile, 5)
        self.assertAlmostEqual(profile["style"]["density"], 1.4, places=4)
        for _ in range(40):
            profiler.learn(profile, 1)
        self.assertGreaterEqual(profile["style"]["density"], 0.6)
        self.assertLessEqual(profile["style"]["density"], 1.4)
        self.assertGreaterEqual(profile["style"]["title_scale"], 0.85)
        self.assertLessEqual(profile["style"]["title_scale"], 1.25)
        self.assertEqual(profile["learn"]["rounds"], 70)
        self.assertEqual(len(profile["learn"]["history"]), 70)

    def test_learned_style_changes_apply(self):
        profile = self._fresh()
        base_size = profiler.apply(profile).title_size
        for _ in range(10):
            profiler.learn(profile, 1)
        self.assertLess(profiler.apply(profile).title_size, base_size)


class OutlineTests(unittest.TestCase):
    def test_outline_parses_as_markdown(self):
        text = outline.build_outline("面向复杂场景的轻量级目标检测方法研究", pages=10)
        doc = md.parse(text)
        self.assertTrue(doc["meta"]["title"])
        slides = doc["slides"]
        self.assertGreaterEqual(len(slides), 4)
        for index, slide in enumerate(slides):
            self.assertTrue(slide["title"].strip(), f"第 {index + 1} 页缺 # 标题")
            self.assertTrue(slide["bullets"], f"第 {index + 1} 页缺要点")
        self.assertTrue(all(slide["note"] for slide in slides))
        self.assertNotIn("```", text)

    def test_kinds_produce_different_structures(self):
        prompt = "Q4 产品规划"
        academic = md.parse(outline.build_outline(prompt, pages=10, kind="academic"))
        business = md.parse(outline.build_outline(prompt, pages=10, kind="business"))
        share = md.parse(outline.build_outline(prompt, pages=10, kind="share"))
        titles = {kind: [s["title"] for s in doc["slides"]]
                  for kind, doc in (("academic", academic), ("business", business),
                                    ("share", share))}
        self.assertEqual(titles["academic"][0], "研究背景与问题")
        self.assertEqual(titles["business"][0], "结论先行")
        self.assertEqual(titles["share"][0], "先问一个问题")
        self.assertNotEqual(titles["academic"], titles["business"])

    def test_kind_aliases_and_auto(self):
        self.assertEqual(outline.resolve_kind("论文答辩"), "academic")
        self.assertEqual(outline.resolve_kind("商业路演"), "business")
        self.assertEqual(outline.resolve_kind("公开分享"), "share")
        self.assertEqual(outline.resolve_kind("auto", "论文答辩 PPT"), "academic")
        self.assertEqual(outline.resolve_kind("auto", "季度经营复盘"), "business")
        self.assertIn(outline.resolve_kind("auto", "随便聊聊"), outline.KINDS)

    def test_density_controls_bullet_count(self):
        low = md.parse(outline.build_outline("人工智能在医疗领域的应用", pages=10,
                                             profile={"style": {"density": 0.6}}))
        high = md.parse(outline.build_outline("人工智能在医疗领域的应用", pages=10,
                                              profile={"style": {"density": 1.4}}))
        low_bullets = sum(len(s["bullets"]) for s in low["slides"])
        high_bullets = sum(len(s["bullets"]) for s in high["slides"])
        self.assertLess(low_bullets, high_bullets)
        for slide in high["slides"]:
            self.assertLessEqual(len(slide["bullets"]), 6)

    def test_density_controls_bullet_length(self):
        low = md.parse(outline.build_outline("人工智能在医疗领域的应用", pages=10,
                                             kind="academic",
                                             profile={"style": {"density": 0.6}}))
        high = md.parse(outline.build_outline("人工智能在医疗领域的应用", pages=10,
                                              kind="academic",
                                              profile={"style": {"density": 1.4}}))
        low_chars = sum(len(b["text"]) for s in low["slides"] for b in s["bullets"])
        high_chars = sum(len(b["text"]) for s in high["slides"] for b in s["bullets"])
        self.assertLess(low_chars, high_chars)
        for doc in (low, high):
            for slide in doc["slides"]:
                for bullet in slide["bullets"]:
                    self.assertGreaterEqual(len(bullet["text"]), 4)

    def test_pages_argument_respected(self):
        for pages in (4, 6, 10, 14):
            doc = md.parse(outline.build_outline("一次分享", pages=pages, kind="share"))
            self.assertGreaterEqual(len(doc["slides"]), 4)
            self.assertLessEqual(len(doc["slides"]), pages + 1)

    def test_explicit_title_wins(self):
        doc = md.parse(outline.build_outline("随便写点东西", pages=6, title="给定标题"))
        self.assertEqual(doc["meta"]["title"], "给定标题")

    def test_empty_prompt_still_works(self):
        doc = md.parse(outline.build_outline("", pages=6))
        self.assertTrue(doc["meta"]["title"])
        self.assertGreaterEqual(len(doc["slides"]), 4)

    def test_markdown_special_chars_are_safe(self):
        doc = md.parse(outline.build_outline("A | B > C # D", pages=6, kind="share"))
        self.assertGreaterEqual(len(doc["slides"]), 4)
        for slide in doc["slides"]:
            self.assertIsNone(slide["table"])

    def test_llm_unavailable_without_env(self):
        with mock.patch.dict(os.environ, {"PPT_LLM_BASE_URL": "", "PPT_LLM_API_KEY": ""}):
            self.assertFalse(outline.llm_available())
            text = outline.build_outline("离线也要能出大纲", pages=6)
        self.assertGreaterEqual(len(md.parse(text)["slides"]), 4)

    def test_llm_failure_falls_back(self):
        env = {
            "PPT_LLM_BASE_URL": "http://127.0.0.1:1/v1",
            "PPT_LLM_API_KEY": "test-key",
            "PPT_LLM_TIMEOUT": "1",
        }
        with mock.patch.dict(os.environ, env):
            self.assertTrue(outline.llm_available())
            text = outline.build_outline("LLM 挂掉时自动降级", pages=8)
        doc = md.parse(text)
        self.assertGreaterEqual(len(doc["slides"]), 4)
        self.assertTrue(doc["meta"]["title"])


class BuildRegressionTests(unittest.TestCase):
    def test_business_minimal_still_9_pages(self):
        with temp_dir() as tmp:
            out = Path(tmp) / "regress.pptx"
            builder.build_file(BUSINESS_MD, out)
            self.assertTrue(out.is_file())
            self.assertEqual(len(Presentation(str(out)).slides), 9)

    def test_build_accepts_theme_obj(self):
        profile = profiler.analyze(EDITORIAL)
        theme = profiler.apply(profile, "editorial")
        text = outline.build_outline("把 AI 当同事之后", pages=8, kind="share",
                                     profile=profile)
        with temp_dir() as tmp:
            md_path = Path(tmp) / "outline.md"
            md_path.write_text(text, encoding="utf-8")
            out = Path(tmp) / "deck.pptx"
            builder.build_file(md_path, out, theme_obj=theme)
            prs = Presentation(str(out))
            self.assertEqual(len(prs.slides), len(md.parse(text)["slides"]) + 1)
            first = prs.slides[0]
            self.assertEqual(tuple(first.background.fill.fore_color.rgb),
                             tuple(theme.bg))

    def test_theme_obj_wins_over_theme_name(self):
        profile = profiler.analyze(EDITORIAL)
        theme = profiler.apply(profile, "editorial")
        text = outline.build_outline("主题对象优先", pages=5)
        with temp_dir() as tmp:
            md_path = Path(tmp) / "o.md"
            md_path.write_text(text, encoding="utf-8")
            out = Path(tmp) / "d.pptx"
            builder.build(md.parse(text), out, theme="minimal", theme_obj=theme)
            first = Presentation(str(out)).slides[0]
            self.assertEqual(tuple(first.background.fill.fore_color.rgb),
                             tuple(theme.bg))
            self.assertNotEqual(tuple(first.background.fill.fore_color.rgb),
                                tuple(themes.get_theme("minimal").bg))


if __name__ == "__main__":
    unittest.main(verbosity=2)
