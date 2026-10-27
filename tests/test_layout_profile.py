"""T15 版式骨架自测：pptx 几何测量 / 图片启发式估计 / 骨架微调不回归。

跑法（工作区根目录）：
    & "$env:USERPROFILE\\.dsh\\dsh-runtimes\\dsh-primary-runtime\\dependencies\\python\\python.exe" tests\\test_layout_profile.py

沙箱不允许写系统临时目录，测试用工作区内的 .tmp/t15-tests/ 当临时区。
"""
from __future__ import annotations

import contextlib
import io
import os
import shutil
import sys
import time
import unittest
import uuid
import zipfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
KIT = ROOT / "kit"
if str(KIT) not in sys.path:
    sys.path.insert(0, str(KIT))

from PIL import Image, ImageDraw                              # noqa: E402
from pptx import Presentation                                 # noqa: E402
from pptx.enum.text import PP_ALIGN                           # noqa: E402
from pptx.util import Inches, Pt                              # noqa: E402

from pptxkit import builder, cli, markdown as md              # noqa: E402
from pptxkit import profile as profiler                       # noqa: E402

TMP_ROOT = ROOT / ".tmp" / "t15-tests"
W, H = 13.333, 7.5

SIMPLE_MD = """---
title: 骨架测试
subtitle: 副标题
author: 张三
org: 测试单位
kicker: 测试
---

[01]
# 章节页

---

# 要点页

- 第一条要点
- 第二条要点

---

# 另一页

- 要点 A
- 要点 B
"""


@contextlib.contextmanager
def temp_dir(name: str = "case"):
    TMP_ROOT.mkdir(parents=True, exist_ok=True)
    path = TMP_ROOT / f"{name}-{os.getpid()}-{uuid.uuid4().hex[:8]}"
    path.mkdir(parents=True, exist_ok=True)
    try:
        yield path
    finally:
        shutil.rmtree(path, ignore_errors=True)


def make_pptx(path: Path) -> None:
    """合成一页：标题在 0.12 页高、左右边距 0.08 页宽、左右两栏正文。"""
    prs = Presentation()
    prs.slide_width = Inches(W)
    prs.slide_height = Inches(H)
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    title_box = slide.shapes.add_textbox(Inches(0.08 * W), Inches(0.12 * H),
                                         Inches(0.84 * W), Inches(0.08 * H))
    run = title_box.text_frame.paragraphs[0].add_run()
    run.text = "合成标题"
    run.font.size = Pt(32)

    for left_ratio in (0.08, 0.56):
        box = slide.shapes.add_textbox(Inches(left_ratio * W), Inches(0.30 * H),
                                       Inches(0.36 * W), Inches(0.55 * H))
        frame = box.text_frame
        for index in range(4):
            paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
            run = paragraph.add_run()
            run.text = f"第 {index + 1} 行正文内容"
            run.font.size = Pt(18)
    prs.save(str(path))


def make_image(path: Path, *, two_columns: bool = False) -> None:
    """白底 + 顶部标题条 + 中部内容块（可选：两块高方差彩色图片区）。"""
    img = Image.new("RGB", (1200, 675), "white")
    draw = ImageDraw.Draw(img)
    # 标题条：x 0.08W–0.5W、y 0.12H–0.20H
    draw.rectangle([96, 81, 600, 135], fill=(35, 35, 35))
    if two_columns:
        for y in range(220, 520, 40):
            draw.rectangle([96, y, 480, y + 12], fill=(70, 70, 70))
            draw.rectangle([720, y, 1104, y + 12], fill=(70, 70, 70))
    else:
        for y in range(220, 380, 30):
            draw.rectangle([96, y, 560, y + 10], fill=(70, 70, 70))
        # 高方差 + 高饱和度图片块：0.5W–0.92W、0.35H–0.80H（随机彩色像素噪声）
        noise = Image.frombytes("RGB", (504, 310), os.urandom(504 * 310 * 3))
        img.paste(noise, (600, 240))
    img.save(path)


def slide_xml(path: Path) -> bytes:
    with zipfile.ZipFile(path) as zf:
        parts = sorted(n for n in zf.namelist()
                       if n.startswith("ppt/slides/slide") and n.endswith(".xml"))
        return b"".join(zf.read(n) for n in parts)


class PptxGeometryTests(unittest.TestCase):
    def test_synthetic_geometry_ratios(self):
        with temp_dir("pptx") as tmp:
            src = tmp / "case.pptx"
            make_pptx(src)
            profile = profiler.analyze(src)
        layout = profile["layout"]
        for key in profiler.LAYOUT_KEYS:
            self.assertIn(key, layout)
        self.assertAlmostEqual(layout["title_top_ratio"], 0.12, delta=0.02)
        self.assertAlmostEqual(layout["title_height_ratio"], 0.08, delta=0.02)
        self.assertAlmostEqual(layout["margin_left_ratio"], 0.08, delta=0.02)
        self.assertAlmostEqual(layout["margin_right_ratio"], 0.08, delta=0.02)
        self.assertEqual(layout["columns"], 2)
        self.assertEqual(layout["alignment"], "left")
        self.assertEqual(layout["source"], "pptx")
        self.assertTrue(0.0 <= layout["image_area_ratio"] <= 1.0)
        self.assertTrue(0.0 <= layout["text_area_ratio"] <= 1.0)

    def test_synthetic_pptx_page_count_and_style_untouched(self):
        with temp_dir("pptx2") as tmp:
            src = tmp / "case.pptx"
            make_pptx(src)
            profile = profiler.analyze(src)
        self.assertEqual(profile["metrics"]["slides"], 1)
        # 几何分析与配色/字体互不影响：没有显式颜色时仍是默认/主题推导值
        self.assertTrue(profile["palette"]["bg"].startswith("#"))
        self.assertTrue(profile["fonts"]["title"])

    def test_bad_pptx_falls_back_to_default_layout(self):
        with temp_dir("bad") as tmp:
            bad = tmp / "broken.pptx"
            bad.write_bytes(b"not a pptx at all")
            profile = profiler.analyze(bad)
        self.assertEqual(profile["layout"], profiler.DEFAULT_LAYOUT)
        self.assertEqual(profile["layout"]["source"], "default")

    def test_fixture_deck_layout(self):
        fixture = ROOT / "tests" / "fixtures" / "case-academic.pptx"
        if not fixture.is_file():
            self.skipTest("缺少测试样例")
        layout = profiler.analyze(fixture)["layout"]
        for key in ("title_top_ratio", "title_height_ratio", "margin_left_ratio",
                    "margin_right_ratio", "margin_top_ratio", "margin_bottom_ratio"):
            self.assertTrue(0.0 <= layout[key] <= 1.0, key)
        self.assertIn(layout["alignment"], ("left", "center"))
        self.assertIn(layout["columns"], (1, 2))


class ImageSkeletonTests(unittest.TestCase):
    def test_synthetic_image_ratios(self):
        with temp_dir("img") as tmp:
            src = tmp / "slide.png"
            make_image(src)
            start = time.time()
            profile = profiler.analyze_image(src)
            elapsed = time.time() - start
        layout = profile["layout"]
        self.assertEqual(layout["source"], "image")
        self.assertTrue(0.0 < layout["confidence"] <= 1.0)
        self.assertAlmostEqual(layout["title_top_ratio"], 0.12, delta=0.03)
        self.assertAlmostEqual(layout["title_height_ratio"], 0.08, delta=0.03)
        self.assertAlmostEqual(layout["margin_left_ratio"], 0.08, delta=0.03)
        self.assertAlmostEqual(layout["margin_right_ratio"], 0.08, delta=0.03)
        self.assertGreater(layout["image_area_ratio"], 0.10)
        self.assertLess(elapsed, 2.0, f"图片分析耗时 {elapsed:.2f}s")

    def test_two_column_image(self):
        with temp_dir("img2") as tmp:
            src = tmp / "columns.png"
            make_image(src, two_columns=True)
            layout = profiler.analyze_image(src)["layout"]
        self.assertEqual(layout["columns"], 2)

    def test_jpg_and_bad_image(self):
        with temp_dir("img3") as tmp:
            png = tmp / "slide.png"
            make_image(png)
            jpg = tmp / "slide.jpg"
            Image.open(png).convert("RGB").save(jpg, quality=92)
            good = profiler.analyze_image(jpg)["layout"]
            broken = tmp / "broken.jpg"
            broken.write_bytes(b"not an image")
            bad = profiler.analyze_image(broken)
        self.assertEqual(good["source"], "image")
        self.assertGreater(good["confidence"], 0.0)
        self.assertEqual(bad["layout"], profiler.DEFAULT_LAYOUT)
        self.assertEqual(bad["metrics"]["slides"], 0)

    def test_image_profile_has_layout_block(self):
        with temp_dir("img4") as tmp:
            src = tmp / "slide.png"
            make_image(src)
            profile = profiler.analyze_image(src, name="图片案例")
        self.assertEqual(profile["name"], "图片案例")
        for key in profiler.LAYOUT_KEYS:
            self.assertIn(key, profile["layout"])
        self.assertEqual(profile["layout"]["source"], "image")


class LayoutHintsTests(unittest.TestCase):
    def test_default_profile_hints_are_neutral(self):
        hints = profiler.layout_hints(profiler._default_profile())
        self.assertEqual(hints["title_shift"], 0.0)
        self.assertEqual(hints["margin_left"], 0.0)
        self.assertEqual(hints["margin_right"], 0.0)

    def test_missing_profile_is_neutral(self):
        self.assertEqual(profiler.layout_hints(None)["title_shift"], 0.0)
        self.assertEqual(profiler.layout_hints({})["margin_right"], 0.0)

    def test_hints_are_clamped(self):
        extreme = {"layout": {**profiler.DEFAULT_LAYOUT,
                              "title_top_ratio": 0.95,
                              "margin_left_ratio": 0.95,
                              "margin_right_ratio": 0.95}}
        hints = profiler.layout_hints(extreme)
        self.assertLessEqual(abs(hints["title_shift"]), 0.15 * H + 1e-6)
        self.assertLessEqual(abs(hints["margin_left"]), 0.10 * W + 1e-6)
        self.assertLessEqual(abs(hints["margin_right"]), 0.10 * W + 1e-6)
        low = {"layout": {**profiler.DEFAULT_LAYOUT, "title_top_ratio": 0.0}}
        self.assertGreaterEqual(profiler.layout_hints(low)["title_shift"],
                                -0.15 * H - 1e-6)

    def test_real_pptx_profile_gives_small_nudges(self):
        fixture = ROOT / "tests" / "fixtures" / "case-academic.pptx"
        if not fixture.is_file():
            self.skipTest("缺少测试样例")
        hints = profiler.layout_hints(profiler.analyze(fixture))
        self.assertLessEqual(abs(hints["title_shift"]), 0.15 * H)
        self.assertEqual(hints["source"], "pptx")


class BuildSkeletonTests(unittest.TestCase):
    def test_no_skeleton_output_is_byte_identical(self):
        with temp_dir("build") as tmp:
            out_a = tmp / "a.pptx"
            out_b = tmp / "b.pptx"
            builder.build(md.parse(SIMPLE_MD), out_a)
            neutral = profiler.layout_hints({"layout": profiler.DEFAULT_LAYOUT})
            builder.build(md.parse(SIMPLE_MD), out_b, skeleton=neutral)
            self.assertEqual(slide_xml(out_a), slide_xml(out_b))

    def test_skeleton_moves_title(self):
        with temp_dir("build2") as tmp:
            out_a = tmp / "a.pptx"
            out_b = tmp / "b.pptx"
            builder.build(md.parse(SIMPLE_MD), out_a)
            hints = dict(profiler.layout_hints({"layout": profiler.DEFAULT_LAYOUT}))
            hints["title_shift"] = 0.5
            hints["margin_left"] = 0.3
            builder.build(md.parse(SIMPLE_MD), out_b, skeleton=hints)
            self.assertNotEqual(slide_xml(out_a), slide_xml(out_b))

    def test_build_file_accepts_skeleton(self):
        with temp_dir("build3") as tmp:
            md_path = tmp / "deck.md"
            md_path.write_text(SIMPLE_MD, encoding="utf-8")
            out = tmp / "deck.pptx"
            hints = dict(profiler.layout_hints({"layout": profiler.DEFAULT_LAYOUT}))
            hints["title_shift"] = -0.2
            builder.build_file(md_path, out, skeleton=hints)
            self.assertTrue(out.is_file())
            self.assertEqual(len(Presentation(str(out)).slides), 4)


class CliProfileTests(unittest.TestCase):
    def test_cli_profile_prints_json_for_pptx_and_image(self):
        with temp_dir("cli") as tmp:
            pptx = tmp / "case.pptx"
            make_pptx(pptx)
            png = tmp / "slide.png"
            make_image(png)
            buffer = io.StringIO()
            with contextlib.redirect_stdout(buffer):
                code_pptx = cli.main(["profile", str(pptx)])
                code_png = cli.main(["profile", str(png)])
            output = buffer.getvalue()
        self.assertEqual(code_pptx, 0)
        self.assertEqual(code_png, 0)
        self.assertIn('"layout"', output)
        self.assertIn('"title_top_ratio"', output)
        self.assertIn('"source": "image"', output)


if __name__ == "__main__":
    unittest.main(verbosity=2)
