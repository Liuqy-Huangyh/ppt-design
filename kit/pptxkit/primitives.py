"""绘图原语：把 python-pptx 的样板代码收敛成几个可靠的小函数。"""
from __future__ import annotations

import math
import re
from pathlib import Path

from lxml import etree
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

from .themes import Theme, rgb

# ---------------------------------------------------------------- 文本

_INLINE = re.compile(r"(\*\*.+?\*\*|`.+?`|\*[^*]+\*)")


def parse_inline(text: str) -> list[tuple[str, dict]]:
    """把 `**粗体**` / `*斜体*` / `` `代码` `` 拆成 (文本, 样式) 片段。"""
    out: list[tuple[str, dict]] = []
    for chunk in _INLINE.split(text):
        if not chunk:
            continue
        if chunk.startswith("**") and chunk.endswith("**") and len(chunk) > 4:
            out.append((chunk[2:-2], {"bold": True}))
        elif chunk.startswith("`") and chunk.endswith("`") and len(chunk) > 2:
            out.append((chunk[1:-1], {"code": True}))
        elif chunk.startswith("*") and chunk.endswith("*") and len(chunk) > 2:
            out.append((chunk[1:-1], {"italic": True}))
        else:
            out.append((chunk, {}))
    return out or [(text, {})]


def _set_cjk_font(run, name: str) -> None:
    """font.name 只写 a:latin，中文字形要单独补 a:ea / a:cs。"""
    rPr = run.font._element
    latin = rPr.find(qn("a:latin"))
    if latin is None:
        latin = etree.SubElement(rPr, qn("a:latin"))
    latin.set("typeface", name)
    anchor = latin
    for tag in ("a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = etree.Element(qn(tag))
            anchor.addnext(el)
        el.set("typeface", name)
        anchor = el


def style_run(run, theme: Theme, *, name=None, size=None, color=None,
              bold=False, italic=False, spacing=None) -> None:
    font = run.font
    font.name = name or theme.font_body
    if size:
        font.size = Pt(size)
    font.bold = bold
    font.italic = italic
    font.color.rgb = color if color is not None else theme.ink
    _set_cjk_font(run, name or theme.font_body)
    if spacing:
        font._element.set("spc", str(int(spacing * 100)))


def _apply_bullet(p, char: str) -> None:
    pPr = p._p.get_or_add_pPr()
    for tag in ("a:buFont", "a:buChar", "a:buNone"):
        old = pPr.find(qn(tag))
        if old is not None:
            pPr.remove(old)
    bu_font = etree.Element(qn("a:buFont"))
    bu_font.set("typeface", "Arial")
    bu = etree.Element(qn("a:buChar"))
    bu.set("char", char)
    anchor = None
    for tag in ("a:tabLst", "a:defRPr"):
        anchor = pPr.find(qn(tag))
        if anchor is not None:
            break
    if anchor is not None:
        anchor.addprevious(bu_font)
        anchor.addprevious(bu)
    else:
        pPr.append(bu_font)
        pPr.append(bu)


def add_para(tf, theme: Theme, text: str = "", *, first=False, size=18,
             color=None, bold=False, italic=False, name=None, font=None,
             align=PP_ALIGN.LEFT, line=1.25, space_before=0, space_after=0,
             bullet=False, bullet_char="•", indent=0.0, hanging=None,
             spacing=None):
    """写一个段落（支持内联粗体/斜体/行内代码）。返回 paragraph。"""
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = align
    p.line_spacing = line
    if space_before:
        p.space_before = Pt(space_before)
    if space_after:
        p.space_after = Pt(space_after)

    if indent or bullet:
        pPr = p._p.get_or_add_pPr()
        marl = int(Inches(indent if indent else 0.30))
        pPr.set("marL", str(marl))
        pPr.set("indent", str(-int(Inches(hanging)) if hanging else -marl))
    if bullet:
        _apply_bullet(p, bullet_char)

    for seg, style in parse_inline(text):
        run = p.add_run()
        run.text = seg
        use_name = "Consolas" if style.get("code") else (font or name)
        style_run(
            run, theme,
            name=use_name,
            size=size,
            color=color,
            bold=bold or style.get("bold", False),
            italic=italic or style.get("italic", False),
            spacing=spacing,
        )
    return p


def textbox(slide, box, *, anchor=MSO_ANCHOR.TOP, wrap=True):
    left, top, width, height = box
    tb = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.auto_size = None
    return tf


# ---------------------------------------------------------------- 形状

def rect(slide, box, *, fill=None, line=None, line_w=1.0, radius=None,
         shape=MSO_SHAPE.RECTANGLE):
    left, top, width, height = box
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else shape
    shp = slide.shapes.add_shape(kind, Inches(left), Inches(top), Inches(width), Inches(height))
    if radius:
        try:
            shp.adjustments[0] = radius
        except (IndexError, ValueError):
            pass
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(line_w)
    shp.shadow.inherit = False
    tf = shp.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    return shp


def hline(slide, left, top, width, color, weight=1.0):
    return rect(slide, (left, top, width, max(weight / 72.0, 0.008)), fill=color)


def vline(slide, left, top, height, color, weight=1.0):
    return rect(slide, (left, top, max(weight / 72.0, 0.008), height), fill=color)


def bg(slide, theme: Theme) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = theme.bg


def notes(slide, text: str) -> None:
    if text and text.strip():
        slide.notes_slide.notes_text_frame.text = text.strip()


# ---------------------------------------------------------------- 图片

def picture(slide, path, box, *, mode="contain", border=None):
    """把图片放进 box 并保持比例；mode=contain 留白居中，cover 裁切铺满。"""
    path = Path(path)
    if not path.exists():
        raise FileNotFoundError(f"图片不存在：{path}")
    left, top, width, height = box
    try:
        from PIL import Image
        with Image.open(path) as im:
            iw, ih = im.size
    except Exception:
        iw, ih = 16, 9
    ratio = (iw / ih) if ih else 16 / 9

    if mode == "cover":
        if width / height > ratio:
            w, h = width, width / ratio
        else:
            h, w = height, height * ratio
        pic = slide.shapes.add_picture(str(path), Inches(left), Inches(top),
                                       Inches(w), Inches(h))
        crop_lr = (w - width) / 2 / w if w > width else 0.0
        crop_tb = (h - height) / 2 / h if h > height else 0.0
        pic.crop_left = pic.crop_right = crop_lr
        pic.crop_top = pic.crop_bottom = crop_tb
        pic.left, pic.top = Inches(left), Inches(top)
        pic.width, pic.height = Inches(width), Inches(height)
    else:
        if width / height > ratio:
            h, w = height, height * ratio
            l, t = left + (width - w) / 2, top
        else:
            w, h = width, width / ratio
            l, t = left, top + (height - h) / 2
        pic = slide.shapes.add_picture(str(path), Inches(l), Inches(t),
                                       Inches(w), Inches(h))
    pic.shadow.inherit = False
    if border is not None:
        pic.line.color.rgb = border
        pic.line.width = Pt(0.75)
    return pic


# ---------------------------------------------------------------- 表格

def _cell_border(cell, color, weight_pt=0.75, sides="tb") -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    for side, tag in (("l", "a:lnL"), ("r", "a:lnR"), ("t", "a:lnT"), ("b", "a:lnB")):
        if side not in sides:
            continue
        old = tcPr.find(qn(tag))
        if old is not None:
            tcPr.remove(old)
    # 以 b,t,r,l 的顺序插到最前面，最终得到 l,r,t,b（schema 顺序）
    for side, tag in (("b", "a:lnB"), ("t", "a:lnT"), ("r", "a:lnR"), ("l", "a:lnL")):
        if side not in sides:
            continue
        ln = etree.Element(qn(tag))
        ln.set("w", str(int(weight_pt * 12700)))
        ln.set("cap", "flat")
        ln.set("cmpd", "sng")
        ln.set("algn", "ctr")
        fill = etree.SubElement(ln, qn("a:solidFill"))
        clr = etree.SubElement(fill, qn("a:srgbClr"))
        clr.set("val", str(color))
        dash = etree.SubElement(ln, qn("a:prstDash"))
        dash.set("val", "solid")
        tcPr.insert(0, ln)


def table(slide, theme: Theme, box, header, rows, *, col_widths=None,
          size=14, header_size=14, row_h=0.42, zebra=True):
    left, top, width, height = box
    n_rows, n_cols = len(rows) + 1, len(header)
    gfx = slide.shapes.add_table(n_rows, n_cols, Inches(left), Inches(top),
                                 Inches(width), Inches(height))
    tbl = gfx.table
    tbl.first_row = False
    tbl.horz_banding = False
    if col_widths:
        total = sum(col_widths)
        for i, w in enumerate(col_widths):
            tbl.columns[i].width = Emu(int(Inches(width) * w / total))
    tbl.rows[0].height = Inches(row_h + 0.06)
    for i in range(1, n_rows):
        tbl.rows[i].height = Inches(row_h)

    for c, text in enumerate(header):
        cell = tbl.cell(0, c)
        cell.fill.solid()
        cell.fill.fore_color.rgb = theme.accent
        cell.margin_left = cell.margin_right = Inches(0.10)
        cell.margin_top = cell.margin_bottom = Inches(0.03)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        cell.text_frame.word_wrap = True
        add_para(cell.text_frame, theme, str(text), first=True, size=header_size,
                 color=rgb("FFFFFF"), bold=True, line=1.0)

    for r, row in enumerate(rows, start=1):
        for c in range(n_cols):
            cell = tbl.cell(r, c)
            cell.fill.solid()
            cell.fill.fore_color.rgb = theme.accent_soft if (zebra and r % 2 == 0) else theme.bg
            cell.margin_left = cell.margin_right = Inches(0.10)
            cell.margin_top = cell.margin_bottom = Inches(0.02)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.text_frame.word_wrap = True
            value = str(row[c]) if c < len(row) else ""
            add_para(cell.text_frame, theme, value, first=True, size=size,
                     color=theme.ink, bold=(c == 0), line=1.05)
            _cell_border(cell, theme.rule, 0.75, sides="b")
    return tbl


# ---------------------------------------------------------------- 度量

def weighted_len(text: str) -> float:
    """中文按 1 字宽、西文按 0.55 字宽估算，用于判断文字是否塞得下。"""
    return sum(1.0 if ord(ch) > 127 else 0.55 for ch in text)


def pick_size(paras: list[str], box_w: float, box_h: float,
              candidates=(26, 24, 22, 21, 20, 19, 18, 17, 16, 15, 14, 13, 12, 11),
              line=1.35, gap=0.4, indent=0.0) -> int:
    """在候选字号里挑一个塞得下的最大值。"""
    if not paras:
        return candidates[0]
    for size in candidates:
        usable_w_pt = max(60.0, (box_w - indent) * 72.0)
        per_line = max(1.0, usable_w_pt / (size * 1.0))
        lines = sum(max(1, math.ceil(weighted_len(p) / per_line)) for p in paras)
        need_h = lines * size * line + (len(paras) - 1) * size * gap
        if need_h <= box_h * 72.0:
            return size
    return candidates[-1]


def slide_number(slide, theme: Theme, index: int, total: int, footer: str = "") -> None:
    """页脚：左下来源/论文题目，右下 03 / 24。"""
    if footer:
        tf = textbox(slide, (0.95, 6.94, 8.6, 0.28))
        add_para(tf, theme, footer, first=True, size=9, color=theme.muted, line=1.0)
    tf = textbox(slide, (11.0, 6.94, 1.38, 0.28))
    add_para(tf, theme, f"{index:02d} / {total:02d}", first=True, size=9,
             color=theme.muted, line=1.0, align=PP_ALIGN.RIGHT)
