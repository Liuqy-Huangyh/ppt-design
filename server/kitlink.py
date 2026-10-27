"""kit/pptxkit 适配层。

契约第 6 节冻结了 `profile.py` / `outline.py` 的签名。这两个模块由 style 负责，
在它们落地之前（或导入失败时）本模块提供**结构一致的最小兜底实现**，保证：

* 服务端代码只调用契约签名，切换到 kit 实现时无需改动；
* 离线、内核缺位时 `/api/profiles`、`/api/generate` 依然可用（T1 验收不阻塞）。

兜底实现不追求与 kit 完全等价的视觉分析质量，只保证字段齐全、取值合法。
"""
from __future__ import annotations

import dataclasses
import datetime as _dt
import importlib
import json
import re
import shutil
import sys
import uuid
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

ROOT = Path(__file__).resolve().parent.parent
KIT_DIR = ROOT / "kit"
if KIT_DIR.is_dir() and str(KIT_DIR) not in sys.path:
    sys.path.insert(0, str(KIT_DIR))

_A = "http://schemas.openxmlformats.org/drawingml/2006/main"
_NS = {"a": _A}

DEFAULT_PALETTE = {
    "bg": "#FFFFFF",
    "ink": "#14213D",
    "muted": "#5C6B82",
    "accent": "#1B3A6B",
    "accent_soft": "#E9EEF7",
    "rule": "#D5DCE8",
}
DEFAULT_FONTS = {"title": "微软雅黑", "body": "微软雅黑"}

FALLBACK_THEME_KEYS = [
    {"key": "academic", "label": "学术答辩"},
    {"key": "editorial", "label": "杂志编辑风"},
    {"key": "minimal", "label": "商务简约"},
]

# 离线大纲兜底用的章节模板（够写出一份能讲的 deck）
_FALLBACK_SECTIONS = [
    ("背景与目标", ["为什么现在要谈这件事", "要解决的核心问题", "本次分享的目标与边界"]),
    ("现状与挑战", ["目前已有的基础", "主要卡点与约束", "可以立刻动手的一步"]),
    ("核心方法", ["方法总览", "关键步骤拆解", "如何衡量是否有效"]),
    ("关键数据", ["口径与样本", "观察到的主要变化", "与预期的差距"]),
    ("实施路径", ["第一阶段：小范围验证", "第二阶段：铺开与协同", "第三阶段：沉淀为流程"]),
    ("风险与对策", ["最可能出问题的地方", "我们的应对策略", "需要谁提供支持"]),
    ("结论与下一步", ["三句话总结", "接下来 30 天要做什么", "联系方式与讨论"]),
]


# --------------------------------------------------------------------------- #
# kit 模块加载
# --------------------------------------------------------------------------- #

def _module(name: str):
    """尝试导入 pptxkit.<name>；缺失或导入失败时返回 None（回退到兜底实现）。"""
    try:
        return importlib.import_module(f"pptxkit.{name}")
    except Exception:
        return None


def using_kit() -> dict:
    """诊断用：报告哪些 kit 模块已被使用。"""
    return {
        "profile": _module("profile") is not None,
        "outline": _module("outline") is not None,
        "kit_dir": str(KIT_DIR),
    }


def theme_keys() -> list[dict]:
    try:
        from pptxkit.themes import THEMES

        return [{"key": key, "label": theme.label} for key, theme in THEMES.items()]
    except Exception:
        return [dict(item) for item in FALLBACK_THEME_KEYS]


def get_theme(name=None):
    """取主题对象；未知主题抛 KeyError（便于 HTTP 层返回 400）。"""
    try:
        from pptxkit.themes import get_theme as _get_theme
    except Exception as exc:  # pragma: no cover - kit 缺失时才会发生
        raise KeyError(f"主题内核不可用：{exc}") from None
    try:
        return _get_theme(name)
    except SystemExit as exc:  # kit 用 SystemExit 报未知主题
        raise KeyError(str(exc) or f"未知主题：{name}") from None


# --------------------------------------------------------------------------- #
# profile.* 适配
# --------------------------------------------------------------------------- #

def analyze(pptx_path, *, name: str = "", source: str = "source.pptx") -> dict:
    mod = _module("profile")
    if mod is not None and hasattr(mod, "analyze"):
        return mod.analyze(pptx_path, name=name, source=source)
    return _analyze_fallback(pptx_path, name=name, source=source)


def save_profile(profile: dict, store_dir) -> dict:
    mod = _module("profile")
    if mod is not None and hasattr(mod, "save"):
        return mod.save(profile, store_dir)
    return _save_fallback(profile, store_dir)


def load_profile(store_dir, profile_id: str) -> dict:
    mod = _module("profile")
    if mod is not None and hasattr(mod, "load"):
        return mod.load(store_dir, profile_id)
    return _load_fallback(store_dir, profile_id)


def list_profiles(store_dir) -> list[dict]:
    mod = _module("profile")
    if mod is not None and hasattr(mod, "list_all"):
        return mod.list_all(store_dir)
    return _list_fallback(store_dir)


def apply_profile(profile: dict, base_theme: str | None = None):
    mod = _module("profile")
    if mod is not None and hasattr(mod, "apply"):
        return mod.apply(profile, base_theme)
    return _apply_fallback(profile, base_theme)


def learn_profile(profile: dict, rating: int, notes: str = "") -> dict:
    mod = _module("profile")
    if mod is not None and hasattr(mod, "learn"):
        return mod.learn(profile, rating, notes)
    return _learn_fallback(profile, rating, notes)


# --------------------------------------------------------------------------- #
# outline.* 适配
# --------------------------------------------------------------------------- #

def build_outline(prompt: str, *, pages: int = 10, kind: str = "auto",
                  profile: dict | None = None, title: str = "") -> str:
    mod = _module("outline")
    if mod is not None and hasattr(mod, "build_outline"):
        return mod.build_outline(prompt, pages=pages, kind=kind,
                                 profile=profile, title=title)
    return local_build_outline(prompt, pages=pages, profile=profile, title=title)


def llm_available() -> bool:
    mod = _module("outline")
    if mod is not None and hasattr(mod, "llm_available"):
        try:
            return bool(mod.llm_available())
        except Exception:
            return False
    import os

    return bool((os.environ.get("PPT_LLM_BASE_URL") or "").strip()
                and (os.environ.get("PPT_LLM_API_KEY") or "").strip())


def local_build_outline(prompt: str, *, pages: int = 10, profile: dict | None = None,
                        title: str = "") -> str:
    """离线启发式大纲（kit.outline 缺位或异常时使用）。"""
    prompt = (prompt or "").strip()
    title = (title or "").strip() or (prompt.splitlines()[0][:40] if prompt else "未命名演示")
    density = 1.0
    if isinstance(profile, dict):
        try:
            density = float((profile.get("style") or {}).get("density", 1.0))
        except (TypeError, ValueError):
            density = 1.0
    per_page = max(2, min(6, int(round(3 * density))))
    body_pages = max(1, max(2, min(40, int(pages or 10))) - 1)  # 封面来自 front matter
    blocks: list[str] = []
    for index in range(body_pages):
        name, bullets = _FALLBACK_SECTIONS[index % len(_FALLBACK_SECTIONS)]
        lines = [f"# {index + 1:02d} {name}"]
        lines.extend(f"- {item}" for item in bullets[:per_page])
        if index == 0:
            if prompt:
                lines.append(f"- 主题：{prompt[:120]}")
            if title:
                lines.append(f"- 本次聚焦：{title[:60]}")
        blocks.append("\n".join(lines))
    return "\n\n".join(blocks) + "\n"


# --------------------------------------------------------------------------- #
# 兜底实现：分析
# --------------------------------------------------------------------------- #

def _now() -> str:
    return _dt.datetime.now().astimezone().isoformat(timespec="seconds")


def _norm_hex(value) -> str | None:
    if not value:
        return None
    text = str(value).strip().lstrip("#")
    if re.fullmatch(r"[0-9A-Fa-f]{6}", text):
        return "#" + text.upper()
    return None


def _mix_white(color: str, ratio: float) -> str:
    """把颜色按 ratio 混向白色（ratio=0 原色，1 纯白）。"""
    try:
        text = color.lstrip("#")
        r, g, b = (int(text[i:i + 2], 16) for i in (0, 2, 4))
    except Exception:
        return color
    ratio = max(0.0, min(1.0, float(ratio)))
    mixed = [int(round(channel + (255 - channel) * ratio)) for channel in (r, g, b)]
    return "#%02X%02X%02X" % tuple(mixed)


def _theme_root(pptx_path) -> ET.Element | None:
    try:
        with zipfile.ZipFile(pptx_path) as archive:
            names = sorted(n for n in archive.namelist()
                           if re.fullmatch(r"ppt/theme/theme\d+\.xml", n))
            if not names:
                return None
            return ET.fromstring(archive.read(names[0]))
    except Exception:
        return None


def _scheme_color(scheme: ET.Element | None, slot: str) -> str | None:
    if scheme is None:
        return None
    node = scheme.find(f"a:{slot}", _NS)
    if node is None:
        return None
    for tag, attr in (("srgbClr", "val"), ("sysClr", "lastClr")):
        child = node.find(f"a:{tag}", _NS)
        if child is not None:
            return _norm_hex(child.get(attr))
    return None


def _palette_from_theme(root: ET.Element | None) -> dict:
    palette = dict(DEFAULT_PALETTE)
    scheme = root.find(".//a:clrScheme", _NS) if root is not None else None
    bg = _scheme_color(scheme, "lt1")
    ink = _scheme_color(scheme, "dk1")
    muted = _scheme_color(scheme, "dk2")
    accent = _scheme_color(scheme, "accent1")
    if bg:
        palette["bg"] = bg
    if ink:
        palette["ink"] = ink
    if muted:
        palette["muted"] = muted
    if accent:
        palette["accent"] = accent
    palette["accent_soft"] = _mix_white(palette["accent"], 0.86)
    palette["rule"] = _mix_white(palette["ink"], 0.82)
    return palette


def _theme_fonts(root: ET.Element | None) -> dict:
    fonts = dict(DEFAULT_FONTS)
    if root is None:
        return fonts
    for key, tag in (("title", "majorFont"), ("body", "minorFont")):
        node = root.find(f".//a:fontScheme/a:{tag}", _NS)
        if node is None:
            continue
        ea = node.find("a:ea", _NS)
        latin = node.find("a:latin", _NS)
        face = ""
        if ea is not None:
            face = (ea.get("typeface") or "").strip()
        if not face and latin is not None:
            face = (latin.get("typeface") or "").strip()
        if face:
            fonts[key] = face
    return fonts


def _median(values: list[float], default: float) -> float:
    if not values:
        return float(default)
    ordered = sorted(values)
    mid = len(ordered) // 2
    if len(ordered) % 2:
        return float(ordered[mid])
    return float((ordered[mid - 1] + ordered[mid]) / 2)


def _analyze_fallback(pptx_path, *, name: str = "", source: str = "source.pptx") -> dict:
    from pptx import Presentation

    prs = Presentation(str(pptx_path))
    slides = list(prs.slides)
    total = len(slides)
    root = _theme_root(pptx_path)
    palette = _palette_from_theme(root)
    fonts = _theme_fonts(root)

    chars = 0
    paragraphs = 0
    title_sizes: list[float] = []
    body_sizes: list[float] = []
    title_fonts: dict[str, int] = {}
    body_fonts: dict[str, int] = {}
    image_slides = 0

    for slide in slides:
        has_image = False
        shapes = list(slide.shapes)
        for shape in shapes:
            try:
                if int(shape.shape_type) == 13:  # MSO_SHAPE_TYPE.PICTURE
                    has_image = True
            except Exception:
                pass
            if not getattr(shape, "has_text_frame", False):
                continue
            try:
                top = float(shape.top) if shape.top is not None else None
            except Exception:
                top = None
            is_title_area = top is not None and top < 1.3 * 914400
            for paragraph in shape.text_frame.paragraphs:
                text = "".join((run.text or "") for run in paragraph.runs).strip()
                if not text:
                    text = (paragraph.text or "").strip()
                if not text:
                    continue
                chars += len(text)
                paragraphs += 1
                sizes = []
                for run in paragraph.runs:
                    if run.font is not None and run.font.size:
                        sizes.append(float(run.font.size.pt))
                    face = run.font.name if run.font is not None else None
                    if face:
                        bucket = title_fonts if is_title_area else body_fonts
                        bucket[face] = bucket.get(face, 0) + 1
                if not sizes:
                    continue
                if is_title_area:
                    title_sizes.append(sizes[0])
                else:
                    body_sizes.append(sizes[0])
        if has_image:
            image_slides += 1

    def _top_font(bucket: dict[str, int], fallback: str) -> str:
        if not bucket:
            return fallback
        return max(bucket.items(), key=lambda kv: kv[1])[0]

    fonts["title"] = _top_font(title_fonts, fonts["title"])
    fonts["body"] = _top_font(body_fonts, fonts["body"])

    try:
        slide_size = [round(prs.slide_width / 914400, 2), round(prs.slide_height / 914400, 2)]
    except Exception:
        slide_size = [13.33, 7.5]

    metrics = {
        "slides": total,
        "chars_per_slide": round(chars / total, 1) if total else 0.0,
        "bullets_per_slide": round(paragraphs / total, 1) if total else 0.0,
        "title_pt": round(_median(title_sizes, 31.0), 1),
        "body_pt": round(_median(body_sizes, 19.0), 1),
        "image_ratio": round(image_slides / total, 2) if total else 0.0,
    }

    return {
        "id": _new_profile_id(),
        "name": (name or "").strip() or "我的案例",
        "source": source or "source.pptx",
        "created_at": _now(),
        "slide_size": slide_size,
        "palette": palette,
        "fonts": fonts,
        "metrics": metrics,
        "style": {"density": 1.0, "title_scale": 1.0, "bullet_char": "▪"},
        "learn": {"rounds": 0, "history": []},
    }


# --------------------------------------------------------------------------- #
# 兜底实现：档案存取 / 主题应用 / 反馈调参
# --------------------------------------------------------------------------- #

def _new_profile_id() -> str:
    return f"pf-{_dt.datetime.now():%Y%m%d}-{uuid.uuid4().hex[:4]}"


def _profiles_dir(store_dir) -> Path:
    return Path(store_dir) / "profiles"


def _cases_dir(store_dir) -> Path:
    return Path(store_dir) / "cases"


def _save_fallback(profile: dict, store_dir) -> dict:
    profile = dict(profile or {})
    profile_id = str(profile.get("id") or _new_profile_id())
    profile["id"] = profile_id
    profiles_dir = _profiles_dir(store_dir)
    cases_dir = _cases_dir(store_dir) / profile_id
    profiles_dir.mkdir(parents=True, exist_ok=True)
    cases_dir.mkdir(parents=True, exist_ok=True)
    source = profile.pop("_source_path", None)
    if source:
        try:
            shutil.copyfile(source, cases_dir / "source.pptx")
            profile["source"] = "source.pptx"
        except OSError:
            pass
    (profiles_dir / f"{profile_id}.json").write_text(
        json.dumps(profile, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return profile


def _load_fallback(store_dir, profile_id: str) -> dict:
    path = _profiles_dir(store_dir) / f"{profile_id}.json"
    if not path.is_file():
        raise FileNotFoundError(f"档案不存在：{profile_id}")
    return json.loads(path.read_text(encoding="utf-8"))


def _list_fallback(store_dir) -> list[dict]:
    items: list[dict] = []
    folder = _profiles_dir(store_dir)
    if not folder.is_dir():
        return items
    for path in folder.glob("*.json"):
        try:
            profile = json.loads(path.read_text(encoding="utf-8"))
        except Exception:
            continue
        if isinstance(profile, dict):
            items.append(profile)
    items.sort(key=lambda p: str(p.get("created_at") or ""), reverse=True)
    return items


def _apply_fallback(profile: dict, base_theme: str | None = None):
    from pptxkit.themes import get_theme as _get_theme, rgb

    base = _get_theme(base_theme)
    palette = dict(profile.get("palette") or {})
    fonts = dict(profile.get("fonts") or {})
    style = dict(profile.get("style") or {})
    try:
        title_scale = float(style.get("title_scale", 1.0))
    except (TypeError, ValueError):
        title_scale = 1.0
    title_scale = min(1.25, max(0.85, title_scale))

    def color(key: str, default):
        value = _norm_hex(palette.get(key))
        return rgb(value) if value else default

    title_size = max(12, int(round(base.title_size * title_scale)))
    return dataclasses.replace(
        base,
        bg=color("bg", base.bg),
        ink=color("ink", base.ink),
        muted=color("muted", base.muted),
        accent=color("accent", base.accent),
        accent_soft=color("accent_soft", base.accent_soft),
        rule=color("rule", base.rule),
        font_title=str(fonts.get("title") or base.font_title),
        font_body=str(fonts.get("body") or base.font_body),
        title_size=title_size,
    )


def _learn_fallback(profile: dict, rating: int, notes: str = "") -> dict:
    style = dict(profile.get("style") or {})
    try:
        density = float(style.get("density", 1.0))
    except (TypeError, ValueError):
        density = 1.0
    try:
        title_scale = float(style.get("title_scale", 1.0))
    except (TypeError, ValueError):
        title_scale = 1.0
    rating = int(rating)
    if rating >= 4:
        density *= 1.03
    elif rating <= 2:
        density *= 0.90
        title_scale *= 0.97
    style["density"] = round(min(1.4, max(0.6, density)), 3)
    style["title_scale"] = round(min(1.25, max(0.85, title_scale)), 3)
    style.setdefault("bullet_char", "▪")
    profile["style"] = style

    learn = dict(profile.get("learn") or {"rounds": 0, "history": []})
    try:
        rounds = int(learn.get("rounds", 0))
    except (TypeError, ValueError):
        rounds = 0
    learn["rounds"] = rounds + 1
    history = list(learn.get("history") or [])
    history.append({
        "at": _now(),
        "rating": rating,
        "density": style["density"],
        "title_scale": style["title_scale"],
        "notes": notes or "",
    })
    learn["history"] = history[-50:]
    profile["learn"] = learn
    return profile
